#!/usr/bin/env python3
"""Build the compact HUD SAFMR asset bundled by the market-rent MVP.

Usage:
  python3 scripts/build_hud_safmr_data.py /path/to/FY27_safmrs.xlsx \
      internal/marketdata/fy2027_selected.json

The source workbook is deliberately not committed. This script preserves the
official FY 2027 40th-percentile gross-rent values for the first supported
Craigslist markets, along with the source metadata needed for a future refresh.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import openpyxl


SOURCE_URL = "https://www.huduser.gov/portal/datasets/fmr/fmr2027/FY27_safmrs.xlsx"
TARGET_AREAS = {
    "Jersey City, NJ HUD Metro FMR Area",
    "Los Angeles-Long Beach-Glendale, CA HUD Metro FMR Area",
    "New York, NY HUD Metro FMR Area",
    "Newark, NJ HUD Metro FMR Area",
    "Oakland-Fremont, CA HUD Metro FMR Area",
    "San Francisco, CA HUD Metro FMR Area",
    "San Jose-Sunnyvale-Santa Clara, CA HUD Metro FMR Area",
}
REQUIRED_HEADERS = {
    "ZIP\nCode": "zip",
    "HUD Fair Market Rent Area Name": "area",
    "SAFMR\n0BR": "studio",
    "SAFMR\n1BR": "one_bedroom",
    "SAFMR\n2BR": "two_bedroom",
    "SAFMR\n3BR": "three_bedroom",
    "SAFMR\n4BR": "four_bedroom",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Official HUD FY 2027 SAFMR workbook")
    parser.add_argument("output", type=Path, help="Destination JSON asset")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook = openpyxl.load_workbook(args.source, read_only=True, data_only=True)
    try:
        worksheet = workbook["SAFMRs"]
        headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        positions = {header: index for index, header in enumerate(headers)}
        missing = sorted(set(REQUIRED_HEADERS) - set(positions))
        if missing:
            raise ValueError(f"HUD workbook schema changed; missing columns: {', '.join(missing)}")

        entries: dict[str, dict[str, object]] = {}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            area = row[positions["HUD Fair Market Rent Area Name"]]
            if area not in TARGET_AREAS:
                continue
            zip_code = str(row[positions["ZIP\nCode"]]).zfill(5)
            rents = [
                row[positions["SAFMR\n0BR"]],
                row[positions["SAFMR\n1BR"]],
                row[positions["SAFMR\n2BR"]],
                row[positions["SAFMR\n3BR"]],
                row[positions["SAFMR\n4BR"]],
            ]
            if not all(isinstance(rent, (int, float)) and rent > 0 for rent in rents):
                raise ValueError(f"invalid HUD SAFMR rents for ZIP code {zip_code}")
            normalized_rents = [int(rent) for rent in rents]
            existing = entries.get(zip_code)
            if existing is None:
                entries[zip_code] = {"areas": {area}, "rents": normalized_rents}
                continue
            if existing["rents"] != normalized_rents:
                raise ValueError(f"conflicting HUD SAFMR rents for ZIP code {zip_code}")
            existing["areas"].add(area)
    finally:
        workbook.close()

    if not entries:
        raise ValueError("no supported-market rows found in HUD workbook")
    for entry in entries.values():
        entry["areas"] = sorted(entry["areas"])
    with args.source.open("rb") as source_file:
        source_sha256 = hashlib.file_digest(source_file, "sha256").hexdigest()
    data = {
        "schema_version": 1,
        "source": {
            "agency": "U.S. Department of Housing and Urban Development",
            "dataset": "FY 2027 Small Area Fair Market Rents",
            "fiscal_year": 2027,
            "effective_date": "2026-10-01",
            "source_url": SOURCE_URL,
            "source_sha256": source_sha256,
            "definition": "40th-percentile gross rent; includes tenant-paid utilities except telephone, cable, satellite, and internet",
            "supported_areas": sorted(TARGET_AREAS),
        },
        "zip_codes": dict(sorted(entries.items())),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, separators=(",", ":"), sort_keys=True) + "\n", encoding="utf-8")
    print(f"wrote {len(entries)} ZIP codes to {args.output}")


if __name__ == "__main__":
    main()
